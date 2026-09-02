"""投标文档引擎 — 文档解析器

支持 PDF (pymupdf) 和 Word (python-docx) 解析。
提取：项目信息、评分标准、技术参数、关键条款、模板段落。
LLM 提取 + regex 回退。
"""

import json
import logging
import re
from pathlib import Path

import httpx

from .models import ParseResult, ProjectInfo, ScoringItem, ScoringType, TechParam, ParamLevel

logger = logging.getLogger("acp-proxy.bidding.parser")

MIMO_BASE = "https://token-plan-cn.xiaomimimo.com/v1"
MIMO_MODEL = "xiaomi/mimo-v2.5-pro"

EXTRACT_PROMPT = """你是招投标文件解析专家。请从以下招标文件内容中提取结构化信息。

返回JSON格式，包含以下字段：
{
  "project_info": {
    "name": "项目名称",
    "budget": "预算金额",
    "deadline": "投标截止时间",
    "method": "采购方式",
    "purchaser": "采购人",
    "agency": "代理机构",
    "location": "项目所在地",
    "bid_open_time": "开标时间",
    "bid_validity": "投标有效期",
    "margin": "保证金",
    "summary": "项目概述（100字以内）"
  },
  "scoring": [
    {"item": "评分项名称", "score": 分值, "type": "价格/技术/服务/资质", "category": "子分类", "criteria": "评分细则", "level_desc": "各档次描述"}
  ],
  "params": [
    {"product": "产品名", "param": "参数名", "value": "要求值", "level": "★/▲/一般", "mandatory": true}
  ],
  "clauses": ["关键条款1原文", "关键条款2原文"]
}

只返回JSON，不要其他文字。

文件内容：
"""


def _extract_text_pdf(file_path: str) -> tuple[str, int]:
    """从PDF提取文本"""
    try:
        import pymupdf
        doc = pymupdf.open(file_path)
        pages = []
        for page in doc:
            pages.append(page.get_text())
        doc.close()
        return "\n".join(pages), len(pages)
    except ImportError:
        logger.warning("pymupdf 未安装，尝试 pdfplumber")
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                pages = [p.extract_text() or "" for p in pdf.pages]
                return "\n".join(pages), len(pages)
        except ImportError:
            raise RuntimeError("需要安装 pymupdf 或 pdfplumber: pip install pymupdf")


def _extract_text_docx(file_path: str) -> tuple[str, int]:
    """从Word文档提取文本"""
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs]
        # 也提取表格内容
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                paragraphs.append(" | ".join(cells))
        return "\n".join(paragraphs), 0
    except ImportError:
        raise RuntimeError("需要安装 python-docx: pip install python-docx")


def _extract_text(file_path: str) -> tuple[str, int, str]:
    """根据文件类型提取文本，返回 (text, page_count, file_type)"""
    p = Path(file_path)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        text, pages = _extract_text_pdf(file_path)
        return text, pages, "pdf"
    elif suffix in (".docx", ".doc"):
        text, _ = _extract_text_docx(file_path)
        return text, 0, "docx"
    else:
        # 尝试当纯文本读
        text = p.read_text(encoding="utf-8", errors="ignore")
        return text, 0, "text"


async def _llm_extract(text: str) -> dict:
    """调用 MiMo API 进行结构化提取"""
    # 截断过长文本（避免token溢出）
    max_chars = 30000
    if len(text) > max_chars:
        text = text[:max_chars] + "\n...(文本已截断)"

    async with httpx.AsyncClient(timeout=120) as client:
        resp = await client.post(
            f"{MIMO_BASE}/chat/completions",
            json={
                "model": MIMO_MODEL,
                "messages": [{"role": "user", "content": EXTRACT_PROMPT + text}],
                "temperature": 0.1,
                "max_tokens": 8000,
            },
        )
        resp.raise_for_status()
        content = resp.json()["choices"][0]["message"]["content"]
        # 提取JSON（可能被markdown包裹）
        json_match = re.search(r'\{[\s\S]*\}', content)
        if json_match:
            return json.loads(json_match.group())
        return json.loads(content)


def _regex_fallback(text: str) -> dict:
    """正则回退提取（当LLM不可用时）"""
    info = {}

    # 项目名称
    for pat in [r'项目名称[：:]\s*(.+?)[\n\r]', r'采购项目[：:]\s*(.+?)[\n\r]']:
        m = re.search(pat, text)
        if m:
            info["name"] = m.group(1).strip()
            break

    # 预算
    for pat in [r'预算[金额]*[：:]\s*([\d,.]+)\s*[万]?元', r'采购预算[：:]\s*([\d,.]+)']:
        m = re.search(pat, text)
        if m:
            info["budget"] = m.group(1).strip()
            break

    # 截止时间
    for pat in [
        r'投标截止[时间日期]*[：:]\s*(.+?)[\n\r]',
        r'递交截止[时间日期]*[：:]\s*(.+?)[\n\r]',
    ]:
        m = re.search(pat, text)
        if m:
            info["deadline"] = m.group(1).strip()
            break

    # 采购方式
    for pat in [r'采购方式[：:]\s*(.+?)[\n\r]', r'招标方式[：:]\s*(.+?)[\n\r]']:
        m = re.search(pat, text)
        if m:
            info["method"] = m.group(1).strip()
            break

    # 评分标准（表格格式）
    scoring = []
    score_pattern = re.finditer(
        r'[（(]?\s*(\d+)\s*[分）)]?\s*[：:-]\s*(.+?)[\n\r]', text
    )
    for m in score_pattern:
        scoring.append(ScoringItem(
            item=m.group(2).strip()[:50],
            score=float(m.group(1)),
            type=ScoringType.TECH,
        ))

    # 技术参数（★/▲标记）
    params = []
    for m in re.finditer(r'[★▲]?\s*(.+?)[：:]\s*(.+?)[\n\r]', text):
        level = ParamLevel.STAR if "★" in m.group(0) else (
            ParamLevel.TRIANGLE if "▲" in m.group(0) else ParamLevel.NORMAL
        )
        params.append(TechParam(
            param=m.group(1).strip()[:30],
            value=m.group(2).strip()[:50],
            level=level,
        ))

    return {
        "project_info": info,
        "scoring": [s.model_dump() for s in scoring[:20]],
        "params": [p.model_dump() for p in params[:50]],
        "clauses": [],
    }


async def parse_document(file_path: str) -> ParseResult:
    """解析招标文件，返回结构化结果

    Args:
        file_path: 文件路径（PDF或Word）

    Returns:
        ParseResult 包含项目信息、评分标准、技术参数等
    """
    text, page_count, file_type = _extract_text(file_path)
    logger.info(f"提取文本完成: {file_path}, {len(text)} chars, {page_count} pages")

    try:
        data = await _llm_extract(text)
        logger.info("LLM提取成功")
    except Exception as e:
        logger.warning(f"LLM提取失败，使用正则回退: {e}")
        data = _regex_fallback(text)

    # 构建 ParseResult
    pi = data.get("project_info", {})
    scoring = [ScoringItem(**s) for s in data.get("scoring", [])]
    params = [TechParam(**p) for p in data.get("params", [])]

    result = ParseResult(
        project_info=ProjectInfo(**pi) if pi else ProjectInfo(),
        scoring=scoring,
        params=params,
        clauses=data.get("clauses", []),
        raw_text=text[:50000],  # 保留前50K字符
        page_count=page_count,
        file_path=file_path,
        file_type=file_type,
    )

    return result
